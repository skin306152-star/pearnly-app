"""Portable Excel pictures and internal navigation without expiring web links."""

from collections import defaultdict
from io import BytesIO

from openpyxl.drawing.image import Image
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Font


def link(cell, sheet_name, row, label):
    cell.value = label
    cell.hyperlink = Hyperlink(
        ref=cell.coordinate, location=f"'{sheet_name}'!A{row}", display=label
    )
    cell.font = Font(color="0563C1", underline="single")


def attach(wb, task, summary, detail):
    images = defaultdict(list)
    for photo in task.get("photos", []):
        images[photo["entry_id"]].append(photo)
    if not images:
        return
    gallery = wb.create_sheet("รูปภาพตรวจนับ")
    gallery.append(["รูปภาพตรวจนับ", "สินค้า / สถานที่ / ผู้ตรวจนับ", "กลับไปสรุป", "กลับไปรายการ"])
    summary_rows = {item["id"]: index for index, item in enumerate(task["items"], 2)}
    detail_rows = {entry.get("id"): index for index, entry in enumerate(task["entries"], 2)}
    active_links = {}
    row_number = 2
    for entry in sorted(task["entries"], key=lambda e: (str(e["item_id"]), e["voided"])):
        group = images[entry.get("id")]
        if not group:
            continue
        start = row_number
        for photo in group:
            values = [
                f'{entry["product_code"]} / {photo["slot"]}',
                f'{entry["product_name"]}\n{entry["warehouse"]} / {entry["location"]}\n'
                f'{entry["counted_by_name"]} / {entry["counted_at"]}\n'
                + ("ยกเลิก (ไม่รวมยอด)" if entry["voided"] else "รวมยอด"),
            ]
            for col, value in enumerate(values, 1):
                cell = gallery.cell(row_number, col, value)
                cell.data_type = "s"
            gallery.row_dimensions[row_number].height = 75
            link(
                gallery.cell(row_number, 3),
                summary.title,
                summary_rows[entry["item_id"]],
                "กลับไปสรุป",
            )
            link(
                gallery.cell(row_number, 4), detail.title, detail_rows[entry["id"]], "กลับไปรายการ"
            )
            picture = Image(BytesIO(photo["content"]))
            scale = min(640 / picture.width, 380 / picture.height, 1)
            picture.width, picture.height = picture.width * scale, picture.height * scale
            gallery.add_image(picture, f"A{row_number + 1}")
            for r in range(row_number + 1, row_number + 22):
                gallery.row_dimensions[r].height = 15
            row_number += 23
        link(
            detail.cell(detail_rows[entry["id"]], 17),
            gallery.title,
            start,
            f"ดูรูปภาพ ({len(group)})",
        )
        if not entry["voided"]:
            first, count = active_links.get(entry["item_id"], (start, 0))
            active_links[entry["item_id"]] = (first, count + len(group))
    for item_id, (row, count) in active_links.items():
        link(summary.cell(summary_rows[item_id], 15), gallery.title, row, f"ดูรูปภาพ ({count})")
