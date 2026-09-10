"""Real database evidence atomicity, tenant boundaries and portable export."""

import base64
from io import BytesIO
from uuid import uuid4
from unittest import TestCase
from unittest import mock

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from core.workspace_context import WorkspaceScope
from services.stocktake import entries, photos, reports, store
from services.stocktake import access
from routes.stocktake_routes import router
from tests.unit import test_stocktake_entries_pg_smoke as scan


def photo(color="red"):
    output = BytesIO()
    Image.new("RGB", (64, 48), color).save(output, "PNG")
    return base64.b64encode(output.getvalue()).decode("ascii")


class StocktakePhotosPgSmoke(TestCase):
    setUpClass = classmethod(scan.StocktakeEntriesPgSmoke.setUpClass.__func__)
    tearDownClass = classmethod(scan.StocktakeEntriesPgSmoke.tearDownClass.__func__)
    setUp = scan.StocktakeEntriesPgSmoke.setUp
    cursor = scan.StocktakeEntriesPgSmoke.cursor
    create = scan.StocktakeEntriesPgSmoke.create
    add = scan.StocktakeEntriesPgSmoke.add
    edit = scan.StocktakeEntriesPgSmoke.edit
    detail = scan.StocktakeEntriesPgSmoke.detail

    def test_http_upload_limits_and_attachment_read_permission(self):
        app = FastAPI()
        app.include_router(router)
        path = f"/api/cowork/stocktakes/{self.task_id}"
        body = dict(request_id=str(uuid4()), quantity="3", warehouse="WH", photos=[photo()] * 6)
        with (
            mock.patch.object(access, "scope_for", return_value=self.scope) as auth,
            TestClient(app) as client,
        ):
            route = path + f'/items/{self.item["id"]}/entries'
            self.assertEqual(client.post(route, json=body).status_code, 422)
            self.assertEqual(self.detail()["entry_total"], 0)
            body["photos"] = [photo()]
            response = client.post(route, json=body)
            self.assertEqual(response.status_code, 200, response.text)
            entry_id = response.json()["entry_id"]
            response = client.get(path + f"/entries/{entry_id}/photos")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers["cache-control"], "private, no-store")
            self.assertEqual(auth.call_args.args[1], "recon.view")
            auth.side_effect = HTTPException(403, detail="authz.forbidden")
            self.assertEqual(client.get(path + f"/entries/{entry_id}/photos").status_code, 403)

    def test_photos_atomic_retry_export_and_delete(self):
        request_id = uuid4()
        receipt = self.add("4", request_id, photo_values=[photo(), photo("blue")])
        self.assertEqual(receipt, self.add("4", request_id, photo_values=[photo(), photo("blue")]))
        entry_id = receipt["entry_id"]
        self.assertEqual(self.detail()["items"][0]["actual_qty"], 4)
        self.assertEqual(self.detail()["entries"][0]["photo_count"], 2)
        self.assertEqual(len(photos.listing(self.scope, self.task_id, entry_id)["photos"]), 2)
        with self.assertRaisesRegex(HTTPException, "operation_conflict"):
            self.add("4", request_id, photo_values=[photo("green")])
        wb = load_workbook(BytesIO(reports.workbook(self.detail())))
        summary, detail, gallery = wb.worksheets
        self.assertEqual(len(gallery._images), 2)
        self.assertEqual(summary["G2"].fill.fgColor.rgb, "00FFF2CC")
        self.assertTrue(summary["O2"].hyperlink.location.endswith("!A2"))
        self.assertTrue(detail["Q2"].hyperlink.location.endswith("!A2"))
        self.assertIsNone(summary["O2"].hyperlink.target)
        wb.close()
        store.delete_task(self.scope, self.task_id)
        with self.cursor() as cur:
            cur.execute(
                "SELECT COUNT(*) AS n FROM cowork_stocktake_photos WHERE entry_id=%s", (entry_id,)
            )
            self.assertEqual(cur.fetchone()["n"], 0)

    def test_limit_rolls_back_quantity_and_edit_preserves_evidence(self):
        receipt = self.add("1", photo_values=[photo()] * 5)
        entry_id = receipt["entry_id"]
        with self.assertRaisesRegex(HTTPException, "photo_limit"):
            entries.write(
                self.scope,
                self.task_id,
                entry_id,
                uuid4(),
                "9",
                "WH",
                "",
                version=0,
                photo_values=[photo()],
            )
        self.assertEqual(self.detail()["items"][0]["actual_qty"], 1)
        self.assertEqual(self.detail()["entries"][0]["version"], 0)
        self.edit(entry_id, "2")
        self.assertEqual(self.detail()["entries"][0]["photo_count"], 5)
        self.edit(entry_id, "2", version=1, voided=True)
        wb = load_workbook(BytesIO(reports.workbook(self.detail())))
        self.assertIsNone(wb.worksheets[0]["O2"].hyperlink)
        self.assertIsNotNone(wb.worksheets[1]["Q2"].hyperlink)
        wb.close()

    def test_photo_isolation_invalid_input_and_closed_task(self):
        entry_id = self.add(photo_values=[photo()])["entry_id"]
        for scope in (
            WorkspaceScope(self.tenant, self.ws + 1, self.user),
            WorkspaceScope(self.other_tenant, self.ws + 2, self.other_user),
        ):
            with self.assertRaises(HTTPException) as caught:
                photos.listing(scope, self.task_id, entry_id)
            self.assertEqual(caught.exception.status_code, 404)
            with self.cursor(
                tenant_id=scope.tenant_id, workspace_client_id=scope.workspace_client_id
            ) as cur:
                cur.execute("SELECT * FROM cowork_stocktake_photos")
                self.assertEqual(cur.fetchall(), [])
        for invalid in ("not base64!", base64.b64encode(b"<svg/>").decode()):
            with self.assertRaisesRegex(HTTPException, "photo_invalid"):
                self.add(photo_values=[invalid])
        self.assertEqual(self.detail()["entry_total"], 1)
        store.close(self.scope, self.task_id)
        with self.assertRaisesRegex(HTTPException, "closed"):
            self.add(photo_values=[photo()])
