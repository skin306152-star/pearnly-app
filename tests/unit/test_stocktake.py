"""Stocktake import and web/LINE permission regression coverage."""

from io import BytesIO
from decimal import Decimal
from unittest import TestCase, mock
from uuid import uuid4

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from starlette.requests import Request

from services.stocktake import access, excel
from routes.stocktake_routes import router


def xlsx(rows, headers=excel.FIELDS):
    wb = Workbook()
    wb.active.append(list(headers))
    for row in rows:
        wb.active.append(row)
    data = BytesIO()
    wb.save(data)
    return data.getvalue()


ROW = ["0012", "สินค้า / 商品", "000012345678", "A", "01", "EA", "2.500001"]


class StocktakeExcel(TestCase):
    def test_precision_identifiers_and_optional_location(self):
        other = ["002", "Other", "QR-002", "", "", "EA", "0"]
        result = excel.parse(xlsx([ROW, other]))
        self.assertEqual(result[0]["book_qty"], Decimal("2.500001"))
        self.assertEqual(result[0]["barcode"], "000012345678")
        self.assertEqual(result[1]["location"], "")

    def test_duplicate_product_rejected_even_in_different_warehouse(self):
        with self.assertRaisesRegex(HTTPException, "duplicate:3"):
            excel.parse(xlsx([ROW, ROW]))
        other = ROW.copy()
        other[3] = "B"
        with self.assertRaisesRegex(HTTPException, "duplicate:3"):
            excel.parse(xlsx([ROW, other]))

    def test_numeric_codes_formulas_and_bad_quantities_rejected(self):
        for column, value in [
            (0, 12),
            (2, 12345678),
            (1, "=1+1"),
            (6, "NaN"),
            (6, "1.0000001"),
            (6, ""),
        ]:
            with self.subTest(column=column, value=value), self.assertRaises(HTTPException):
                row = ROW.copy()
                row[column] = value
                excel.parse(xlsx([row]))
        for quantity in ("-1", "Infinity", "", "1e99"):
            with self.subTest(quantity=quantity), self.assertRaises(HTTPException):
                excel.quantity(quantity, nonnegative=True)
        self.assertEqual(excel.quantity("0", nonnegative=True), Decimal(0))

    def test_invalid_empty_and_oversize(self):
        for data in (b"not excel", xlsx([]), b"0" * (excel.MAX_BYTES + 1)):
            with self.assertRaises(HTTPException):
                excel.parse(data)
        with self.assertRaisesRegex(HTTPException, "headers_invalid"):
            excel.parse(xlsx([ROW], headers=["wrong"]))

    def test_export_preserves_null_zero_and_formula_safety(self):
        item = dict(zip(excel.FIELDS, ROW))
        item.update(
            product_name='=HYPERLINK("https://example.invalid")',
            book_qty=Decimal("2.500001"),
            actual_qty=None,
            counted_at=None,
            counted_by=None,
        )
        wb = load_workbook(BytesIO(excel.workbook([item, {**item, "actual_qty": Decimal(0)}])))
        ws = wb.active
        self.assertIsNone(ws["H2"].value)
        self.assertIsNone(ws["I2"].value)
        self.assertEqual(ws["H3"].value, 0)
        self.assertEqual(ws["I3"].value, -2.500001)
        self.assertEqual(ws["B2"].data_type, "s")
        wb.close()

    def test_export_localized_headers_and_large_quantity_precision(self):
        item = dict(zip(excel.FIELDS, ROW))
        item.update(
            book_qty=Decimal("12345678901234.123456"),
            actual_qty=None,
            counted_at=None,
            counted_by=None,
        )
        for lang in ("zh", "th", "en", "ja"):
            wb = load_workbook(BytesIO(excel.workbook([item], lang=lang)))
            self.assertEqual(wb.active["A1"].value, excel.LABELS[lang][0])
            self.assertEqual(wb.active["G2"].value, "12345678901234.123456")
            wb.close()

    def test_template_can_be_filled_without_header_changes(self):
        wb = load_workbook(BytesIO(excel.workbook()))
        self.assertEqual([cell.value for cell in wb.active[1]], excel.TEMPLATE_HEADERS)
        item = dict(zip(excel.FIELDS, ROW))
        for index, value in enumerate((item[k] for k in excel.TEMPLATE_FIELDS), 1):
            wb.active.cell(2, index, value)
        data = BytesIO()
        wb.save(data)
        self.assertEqual(len(excel.parse(data.getvalue())), 1)

    def test_five_column_template_and_internal_qr_match_exactly(self):
        row = ["0001", "QR product", "Company/QR?ID=0001", "EA", "10"]
        for headers in (excel.TEMPLATE_HEADERS[:5], excel.TEMPLATE_FIELDS[:5]):
            result = excel.parse(xlsx([row], headers))[0]
            self.assertEqual(result["barcode"], "Company/QR?ID=0001")
            self.assertEqual((result["warehouse"], result["location"]), ("", ""))
        for second in (
            ["0002", "Other", row[2], "EA", "5"],
            [row[2], "Other", "Distinct QR", "EA", "5"],
        ):
            with self.assertRaisesRegex(HTTPException, "ambiguous_code"):
                excel.parse(xlsx([row, second], excel.TEMPLATE_FIELDS[:5]))

    def test_export_http_uses_thai_even_with_legacy_page_language(self):
        app = FastAPI()
        app.include_router(router)
        client = TestClient(app)
        item = dict(zip(excel.FIELDS, ROW))
        item.update(
            book_qty=Decimal("10"), actual_qty=Decimal("9"), counted_at=None, counted_by=None
        )
        with (
            mock.patch("routes.stocktake_routes.stocktake_access.scope_for", return_value={}),
            mock.patch("routes.stocktake_routes.store.detail", return_value={"items": [item]}),
        ):
            for query in ("", "?lang=zh", "?lang=en", "?lang=ja"):
                with self.subTest(query=query):
                    response = client.get(f"/api/cowork/stocktakes/{uuid4()}/export{query}")
                    self.assertEqual(response.status_code, 200)
                    wb = load_workbook(BytesIO(response.content))
                    self.assertEqual([cell.value for cell in wb.active[1]], excel.LABELS["th"])
                    self.assertEqual(wb.active["I2"].value, -1)
                    wb.close()


class StocktakeAccess(TestCase):
    def request(self, headers=()):
        return Request({"type": "http", "headers": headers})

    def test_web_uses_cowork_permissions_and_requires_company(self):
        user = {"id": "u", "tenant_id": "t", "entry": "cowork"}
        with mock.patch.object(access, "require_perm", return_value=user) as perm:
            with self.assertRaisesRegex(HTTPException, "workspace.required"):
                access.scope_for(self.request(), "recon.create")
            perm.assert_called_once()
        for entry in ("pos", "erp", "ai", "dms"):
            with (
                mock.patch.object(access, "require_perm", return_value={**user, "entry": entry}),
                self.assertRaises(HTTPException),
            ):
                access.authorize(self.request(), "recon.view")

    def test_assigned_company_scope_is_enforced(self):
        user = {"id": "u", "tenant_id": "t", "entry": "cowork"}
        request = self.request([(b"x-workspace-client-id", b"42")])
        with (
            mock.patch.object(access, "require_perm", return_value=user),
            mock.patch.object(access, "require_workspace_id", return_value=42),
            mock.patch.object(
                access, "check_workspace_scope", side_effect=HTTPException(404)
            ) as scope,
        ):
            with self.assertRaises(HTTPException) as caught:
                access.scope_for(request, "recon.view")
            self.assertEqual(caught.exception.status_code, 404)
            scope.assert_called_once_with(request, user, 42)

    def test_line_token_cannot_be_replaced_by_unverified_id(self):
        request = self.request([(b"x-stocktake-line", b"1"), (b"authorization", b"Bearer forged")])
        with (
            mock.patch.object(access, "_secret", return_value="test-secret-long-enough-for-tests"),
            self.assertRaises(HTTPException) as caught,
        ):
            access.authorize(request, "recon.view")
        self.assertEqual(caught.exception.status_code, 401)

    def test_revoked_line_membership_rejected(self):
        import jwt

        token = jwt.encode(
            {"sub": "line", "membership": "old", "aud": "cowork_stocktake"},
            "test-secret-long-enough-for-tests",
            algorithm=access.JWT_ALGORITHM,
        )
        request = self.request(
            [(b"x-stocktake-line", b"1"), (b"authorization", ("Bearer " + token).encode())]
        )
        with (
            mock.patch.object(access, "_secret", return_value="test-secret-long-enough-for-tests"),
            mock.patch.object(access, "resolve_active_identity", return_value=None),
            self.assertRaises(HTTPException),
        ):
            access.authorize(request, "recon.view")

    def test_public_login_verifies_existing_cowork_channel(self):
        with (
            mock.patch.object(access, "verify_id_token", return_value=None) as verify,
            mock.patch.object(access, "resolve_active_identity", return_value=None),
            self.assertRaises(HTTPException),
        ):
            access.line_login("invalid")
        verify.assert_called_once_with("invalid", "LINE_COWORK_LIFF_ID")

    def test_http_routes_reject_without_auth_and_use_threadpool(self):
        app = FastAPI()
        app.include_router(router)
        with (
            TestClient(app) as client,
            mock.patch.object(access, "require_perm", side_effect=HTTPException(401)),
        ):
            self.assertEqual(client.get("/api/cowork/stocktakes").status_code, 401)
            self.assertEqual(
                client.post(
                    "/api/cowork/stocktakes",
                    data={"name": "x", "request_id": str(uuid4())},
                    files={"file": ("test.xlsx", xlsx([ROW]))},
                ).status_code,
                401,
            )
        import inspect

        self.assertTrue(
            all(
                not inspect.iscoroutinefunction(r.endpoint)
                for r in app.routes
                if r.path.startswith("/api/cowork/stocktakes")
            )
        )

    def test_menu_keeps_existing_erp_action_and_adds_jump(self):
        from services.cowork_line.menu_cards import menu_card, ACTION_ERP_START
        import json

        with mock.patch.dict("os.environ", {"LINE_COWORK_LIFF_ID": "123-cowork"}):
            payload = json.dumps(menu_card("zh"))
        self.assertIn(ACTION_ERP_START, payload)
        self.assertIn("https://liff.line.me/123-cowork?flow=cowork-stocktake&draft=list", payload)
