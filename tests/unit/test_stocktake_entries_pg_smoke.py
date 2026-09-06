"""Real PostgreSQL entry receipts, aggregate corrections, isolation and HTTP export."""

from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest import TestCase, mock
from uuid import uuid4

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.workspace_context import WorkspaceScope
from routes.stocktake_routes import router
from services.stocktake import access, entries, excel, store
from tests.unit.test_stocktake import ROW, xlsx
from tests.unit import test_stocktake_pg_smoke as legacy


class StocktakeEntriesPgSmoke(TestCase):
    setUpClass = classmethod(legacy.StocktakePgSmoke.setUpClass.__func__)
    tearDownClass = classmethod(legacy.StocktakePgSmoke.tearDownClass.__func__)
    cursor = legacy.StocktakePgSmoke.cursor
    setUp = legacy.StocktakePgSmoke.setUp

    def create(self):
        return store.create(self.scope, "Scan count", excel.parse(xlsx([ROW])), uuid4(), "scan")[
            "id"
        ]

    def add(self, qty="1", request_id=None, scope=None, **kwargs):
        return entries.write(
            scope or self.scope,
            self.task_id,
            self.item["id"],
            request_id or uuid4(),
            Decimal(qty),
            "Free WH",
            "Free location",
            **kwargs,
        )

    def edit(self, entry_id, qty="1", version=0, voided=False):
        return entries.write(
            self.scope,
            self.task_id,
            entry_id,
            uuid4(),
            Decimal(qty),
            "Other WH",
            "Other location",
            version=version,
            voided=voided,
        )

    def detail(self):
        return store.detail(self.scope, self.task_id, export=True)

    def test_receipts_sum_zero_null_and_correction_audit(self):
        self.assertIsNone(self.detail()["items"][0]["actual_qty"])
        operation = uuid4()
        first = self.add("0", operation)
        self.assertEqual(first, self.add("0.000000", operation))
        self.assertEqual(self.detail()["items"][0]["actual_qty"], 0)
        with self.assertRaisesRegex(HTTPException, "operation_conflict"):
            self.add("1", operation)
        second = self.add("2")
        third = self.add("2")
        self.assertEqual(self.detail()["items"][0]["actual_qty"], 4)
        self.edit(second["entry_id"], "3")
        with self.assertRaisesRegex(HTTPException, "conflict"):
            self.edit(second["entry_id"], "4")
        self.edit(third["entry_id"], "2", voided=True)
        self.assertEqual(self.detail()["items"][0]["actual_qty"], 3)
        self.edit(first["entry_id"], "0", voided=True)
        self.edit(second["entry_id"], "3", version=1, voided=True)
        self.assertIsNone(self.detail()["items"][0]["actual_qty"])
        with self.cursor() as cur:
            cur.execute(
                "SELECT previous,current FROM cowork_stocktake_entry_events WHERE entry_id=%s ORDER BY version",
                (second["entry_id"],),
            )
            audit = cur.fetchall()
        self.assertEqual(len(audit), 3)
        self.assertEqual(Decimal(audit[1]["previous"]["quantity"]), 2)
        self.assertEqual(Decimal(audit[1]["current"]["quantity"]), 3)
        self.assertTrue(audit[2]["current"]["voided"])

    def test_entry_isolation_and_direct_rls(self):
        entry = self.add()["entry_id"]
        for scope in (
            WorkspaceScope(self.tenant, self.ws + 1, self.user),
            WorkspaceScope(self.other_tenant, self.ws + 2, self.other_user),
        ):
            for action in (
                lambda: self.add(scope=scope),
                lambda: entries.listing(scope, self.task_id, 50, 0),
                lambda: store.detail(scope, self.task_id, export=True),
                lambda: entries.write(
                    scope, self.task_id, entry, uuid4(), Decimal(2), "WH", "", version=0
                ),
            ):
                with self.assertRaises(HTTPException) as caught:
                    action()
                self.assertEqual(caught.exception.status_code, 404)
            with self.cursor(
                tenant_id=scope.tenant_id, workspace_client_id=scope.workspace_client_id
            ) as cur:
                for table in ("cowork_stocktake_entries", "cowork_stocktake_entry_events"):
                    cur.execute(f"SELECT * FROM {table}")
                    self.assertEqual(cur.fetchall(), [])

    def race(self, actions):
        barrier = Barrier(len(actions))

        def run(action):
            barrier.wait()
            try:
                action()
                return 200
            except HTTPException as exc:
                return exc.status_code

        with ThreadPoolExecutor(len(actions)) as executor:
            return list(executor.map(run, actions))

    def test_concurrent_adds_and_same_receipt_have_exactly_one_effect_each(self):
        self.assertEqual(self.race([lambda: self.add("2"), lambda: self.add("3")]), [200, 200])
        operation = uuid4()
        self.assertEqual(self.race([lambda: self.add("4", operation)] * 2), [200, 200])
        detail = self.detail()
        self.assertEqual(detail["items"][0]["actual_qty"], 9)
        self.assertEqual(detail["entry_total"], 3)

    def test_concurrent_edits_and_close_are_serialized(self):
        entry = self.add()["entry_id"]
        self.assertEqual(
            sorted(self.race([lambda: self.edit(entry, "2"), lambda: self.edit(entry, "3")])),
            [200, 409],
        )
        result = self.race([lambda: store.close(self.scope, self.task_id), lambda: self.add("4")])
        self.assertEqual(result[0], 200)
        self.assertIn(result[1], (200, 409))
        self.assertEqual(self.detail()["status"], "closed")
        with self.assertRaisesRegex(HTTPException, "closed"):
            self.add()

    def test_successful_retry_after_close_and_overflow_rollback(self):
        operation = uuid4()
        receipt = self.add("99999999999999.999999", operation)
        with self.assertRaises(HTTPException):
            self.add("0.000001")
        self.assertEqual(self.detail()["entry_total"], 1)
        store.close(self.scope, self.task_id)
        self.assertEqual(self.add("99999999999999.999999", operation), receipt)

    def test_old_write_endpoint_cannot_overwrite_scan_totals(self):
        self.add("3")
        with self.assertRaisesRegex(HTTPException, "refresh_required"):
            store.count(self.scope, self.task_id, self.item["id"], Decimal(1), 0)
        self.assertEqual(self.detail()["items"][0]["actual_qty"], 3)

    def test_http_import_add_edit_and_two_sheet_thai_export(self):
        app = FastAPI()
        app.include_router(router)
        with (
            TestClient(app) as client,
            mock.patch.object(access, "scope_for", return_value=self.scope),
        ):
            response = client.post(
                "/api/cowork/stocktakes",
                data={"name": "Scan HTTP", "request_id": str(uuid4())},
                files={"file": ("test.xlsx", xlsx([ROW]))},
            )
            self.assertEqual(response.status_code, 200, response.text)
            path = "/api/cowork/stocktakes/" + response.json()["id"]
            detail = client.get(path).json()
            self.assertEqual(detail["count_mode"], "scan")
            body = {
                "request_id": str(uuid4()),
                "quantity": "0",
                "warehouse": "New warehouse",
                "location": "Rack X",
            }
            response = client.post(
                path + "/items/" + detail["items"][0]["id"] + "/entries", json=body
            )
            self.assertEqual(response.status_code, 200, response.text)
            entry_id = response.json()["entry_id"]
            body.update(request_id=str(uuid4()), quantity="2.5", version=0, voided=False)
            response = client.patch(path + "/entries/" + entry_id, json=body)
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(client.get(path + "/entries?limit=1").json()["total"], 1)
            wb = load_workbook(BytesIO(client.get(path + "/export?lang=en").content))
            self.assertEqual(wb.sheetnames, ["สรุปผลตรวจนับ", "รายละเอียดการนับ"])
            self.assertEqual(wb.worksheets[0]["F2"].value, 2.5)
            self.assertEqual(wb.worksheets[1]["D2"].value, "New warehouse")
            self.assertEqual(wb.worksheets[1]["G2"].value, 2.5)
            wb.close()
