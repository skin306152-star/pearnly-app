# -*- coding: utf-8 -*-
"""Credits 计费 · 定价策略 + 成本/单位估算(无 DB · 无扣费)

从 db.py 抽出(REFACTOR-B2 · 纯搬家 · 0 逻辑改)。
PDF 分段阶梯价 + Excel/Word/CSV 按字符计价。db.charge_ocr 经 db.py 尾 re-export 调本模块
estimate_*(裸名解析到 db 命名空间的 re-export)· app/recon_*/services.ocr 走 db.estimate_* 不变。

计费单位判据(2026-08-13 收口)也归这里单源:扩展名分类(EXCEL_BILLING_EXTS)、
字符估算(excel_char_count_estimate)、页折算(pdf_billing_units)、预检/事后单位
汇总(estimate_recon_units / billed_units_for_parses)。此前同一套判据抄了 5 份散在
routes 与 recon_jobs,预检与实扣口径一漂就是在钱上撒谎。
"""

import logging
from decimal import Decimal as _DecV21, ROUND_HALF_UP as _RH_V21
from typing import Iterable, List, Tuple
import math as _math_v21

logger = logging.getLogger(__name__)

PDF_TIER1_LIMIT_V21 = 200
PDF_TIER1_PRICE_V21 = _DecV21("1.50")
PDF_TIER2_PRICE_V21 = _DecV21("0.75")
EXCEL_CHARS_PER_SATANG_V21 = 50
EXCEL_SATANG_PRICE_V21 = _DecV21("0.01")


def estimate_pdf_cost_thb(pages_used_this_month: int, page_count: int) -> _DecV21:
    """估算 PDF N 页的总成本 · 跨界自动拆段
    v0.21 改: 调用端传 pages_used_this_month · 不再查 DB · 与前置 combined 查询复用
    """
    n = max(0, int(page_count or 0))
    if n == 0:
        return _DecV21("0.00")
    used = max(0, int(pages_used_this_month or 0))
    tier1_remaining = max(0, PDF_TIER1_LIMIT_V21 - used)
    tier1_pages = min(n, tier1_remaining)
    tier2_pages = n - tier1_pages
    cost = (PDF_TIER1_PRICE_V21 * tier1_pages) + (PDF_TIER2_PRICE_V21 * tier2_pages)
    return cost.quantize(_DecV21("0.01"), rounding=_RH_V21)


def estimate_excel_cost_thb(char_count: int) -> _DecV21:
    """Excel/Word/CSV 按字符计费 · 50 字符 = 1 satang · 向上取整"""
    n = max(0, int(char_count or 0))
    if n == 0:
        return _DecV21("0.00")
    satang = _math_v21.ceil(n / EXCEL_CHARS_PER_SATANG_V21)
    return (EXCEL_SATANG_PRICE_V21 * satang).quantize(_DecV21("0.01"), rounding=_RH_V21)


# 订阅套餐目录:额度(张/周期)· 月费 · 超额单价(超出额度后每张从余额扣)。
# 周期 = 订阅日起 SUBSCRIPTION_CYCLE_DAYS 天 · 到期自动从余额续订(见 services/billing/subscription.py)。
SUBSCRIPTION_PLANS = {
    "S": {"quota": 100, "fee": _DecV21("150"), "over_rate": _DecV21("1.50")},
    "M": {"quota": 200, "fee": _DecV21("250"), "over_rate": _DecV21("1.25")},
    "L": {"quota": 500, "fee": _DecV21("500"), "over_rate": _DecV21("1.00")},
}
SUBSCRIPTION_CYCLE_DAYS = 30

# 文档(Excel/Word/CSV)按字符计费 · 折算成套餐「张」额度的基准价:
# 按量成本 ÷ DOC_QUOTA_REF_PRICE 向上取整 = 该文档占用的额度张数(与按量一档页价 ฿1.50 对齐 · 跨套餐一致)。
DOC_QUOTA_REF_PRICE = _DecV21("1.50")


def subscription_plan_spec(plan_code: str) -> dict | None:
    """套餐码(S/M/L · 大小写不敏感)→ {quota, fee, over_rate};未知码返 None。"""
    return SUBSCRIPTION_PLANS.get((plan_code or "").strip().upper())


def doc_quota_pages(char_count: int) -> int:
    """Excel/Word/CSV 文档折算成套餐额度张数:按量成本 ÷ ฿1.50 · 向上取整。

    PDF/图片按物理页数直接占额度(1 页 = 1 张),无需本函数;字符计费文档单位不一致才折算。
    """
    cost = estimate_excel_cost_thb(char_count)
    if cost <= 0:
        return 0
    return _math_v21.ceil(cost / DOC_QUOTA_REF_PRICE)


def estimate_recon_cost_thb(
    pages_used_this_month: int, pdf_units: int, excel_chars: int = 0
) -> _DecV21:
    """一批文件的预检总估价 = PDF 阶梯价 + Excel 字符折算成张 × 基准张价。

    Excel/CSV 字符文件按 doc_quota_pages 折算成额度张数、每张按 DOC_QUOTA_REF_PRICE
    计(与按量一档页价对齐)。估值口径 = 解析前文本量,与事后实扣的解析后字符数可能
    有小偏差;预检从宽,不加安全系数。
    """
    pdf_cost = estimate_pdf_cost_thb(pages_used_this_month, pdf_units)
    excel_cost = _DecV21(doc_quota_pages(excel_chars)) * DOC_QUOTA_REF_PRICE
    return (pdf_cost + excel_cost).quantize(_DecV21("0.01"), rounding=_RH_V21)


# ── 计费单位判据(预检估价与事后实扣的共同单源)──────────────────────

# 「字符」档扩展名:预检与 charge_ocr_async 的分类必须逐字同一组,否则同一份文件
# 预检按页、实扣按字符(或反之),402 拦截与真实扣费对不上。
EXCEL_BILLING_EXTS = frozenset({".xlsx", ".xls", ".xlsm", ".csv", ".tsv", ".txt", ".docx", ".doc"})

ROWS_PER_PAGE_BILLING = 40  # 居中计费:一页约 40 笔 · 防密集账单按页低估


def file_ext(filename: str) -> str:
    """小写扩展名(带点)· 无扩展名返 ''。"""
    return ("." + (filename or "").lower().rsplit(".", 1)[-1]) if "." in (filename or "") else ""


def pdf_billing_units(page_count: int, row_count: int) -> int:
    """银行对账 PDF/图片计费『页数』· 居中口径 max(实际页数, ⌈行数/40⌉)。

    对齐 ฿1.5/页规则 · v118.35.0.58 修复此前误按交易行数计费(超收 10-34 倍)的 bug。
    既不让多页大账单超收 · 也不让一页塞很多笔的密集账单被低估 · 图片=1 页。
    """
    pages = max(1, int(page_count or 0))
    rows = max(0, int(row_count or 0))
    return max(pages, _math_v21.ceil(rows / ROWS_PER_PAGE_BILLING))


def excel_char_count_estimate(file_bytes: bytes, filename: str) -> int:
    """估算 Excel/CSV/Word 文件的总字符数(计费 units)· 读不出时粗估降级。"""
    if not file_bytes:
        return 0
    fn = (filename or "").lower()
    try:
        if fn.endswith(".xlsx") or fn.endswith(".xlsm") or fn.endswith(".xls"):
            try:
                import openpyxl
                import io

                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
                total = 0
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        for c in row:
                            if c is not None:
                                total += len(str(c))
                return total
            except Exception:
                return max(0, len(file_bytes) // 4)  # 粗估降级
        elif fn.endswith(".csv") or fn.endswith(".tsv") or fn.endswith(".txt"):
            try:
                return len(file_bytes.decode("utf-8", errors="ignore"))
            except Exception:
                return 0
        elif fn.endswith(".docx") or fn.endswith(".doc"):
            try:
                import docx
                import io

                doc = docx.Document(io.BytesIO(file_bytes))
                return sum(len(p.text) for p in doc.paragraphs)
            except Exception:
                return max(0, len(file_bytes) // 2)
    except Exception as e:
        logger.warning(f"excel_char_count_estimate error fn={fn}: {e}")
    return 0


def estimate_recon_units(files: List[Tuple[bytes, str]]) -> Tuple[int, int]:
    """一批 (bytes, filename) 的预检计费单位 (pdf_units, excel_chars)。

    Excel/CSV/Word 走便宜的结构读取估字符;其余按物理页数计(图片/读不出页数的
    损坏件按 1 页)。预检拿不到解析行数,物理页数是事后实扣 pdf_billing_units
    (页与行折算取大)的下限 —— 多页 PDF 不再按「1 件 1 页」低估打穿余额。
    """
    from services.ocr.pdf_utils import count_pdf_pages

    pdf_units = 0
    excel_chars = 0
    for b, fn in files or []:
        if file_ext(fn) in EXCEL_BILLING_EXTS:
            excel_chars += int(excel_char_count_estimate(b, fn) or 0)
        else:
            pdf_units += max(1, int(count_pdf_pages(b) or 0))
    return pdf_units, excel_chars


def billed_units_for_parses(pairs: Iterable) -> Tuple[int, int]:
    """成功解析件的事后扣费单位 (pdf_units, excel_chars)。

    pairs = [(解析结果 dict, (bytes, filename)), …]。失败件与 0 行件不收钱(全站
    「失败不收钱」口径);字符档按估算字符,其余按 pdf_billing_units。四个对账入口
    (同步两路 + worker 两路)共用此判据。
    """
    from services.ocr.pdf_utils import count_pdf_pages

    pdf_units = 0
    excel_chars = 0
    for r, (b, fn) in pairs or []:
        if not r.get("ok"):
            continue
        row_count = len(r.get("rows") or [])
        if row_count == 0:
            continue
        if file_ext(fn) in EXCEL_BILLING_EXTS:
            excel_chars += int(excel_char_count_estimate(b, fn) or 0)
        else:
            pdf_units += pdf_billing_units(count_pdf_pages(b) or 1, row_count)
    return pdf_units, excel_chars
