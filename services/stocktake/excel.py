"""Bounded XLSX import, exact quantities, and formula-safe export."""

from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from xml.etree.ElementTree import ParseError

FIELDS = ("product_code", "product_name", "barcode", "warehouse", "location", "unit", "book_qty")
LABELS = {
    "zh": [
        "商品编号",
        "商品名称",
        "条码",
        "仓库",
        "库位",
        "单位",
        "账面数量",
        "实盘数量",
        "差异",
        "盘点时间（UTC）",
        "盘点人 ID",
    ],
    "en": [
        "Product code",
        "Product name",
        "Barcode",
        "Warehouse",
        "Location",
        "Unit",
        "Book quantity",
        "Counted quantity",
        "Difference",
        "Counted at (UTC)",
        "Counted by ID",
    ],
    "th": [
        "รหัสสินค้า",
        "ชื่อสินค้า",
        "บาร์โค้ด",
        "คลังสินค้า",
        "ช่องเก็บ",
        "หน่วย",
        "จำนวนตามบัญชี",
        "จำนวนที่นับได้",
        "ผลต่าง",
        "เวลาที่นับ (UTC)",
        "รหัสผู้ตรวจนับ",
    ],
    "ja": [
        "商品コード",
        "商品名",
        "バーコード",
        "倉庫",
        "棚番",
        "単位",
        "帳簿数量",
        "実棚数量",
        "差異",
        "棚卸日時（UTC）",
        "担当者ID",
    ],
}

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000


def quantity(value, *, nonnegative=False):
    try:
        if isinstance(value, bool) or value is None or not str(value).strip():
            raise ValueError()
        result = Decimal(str(value).strip())
        if not result.is_finite() or abs(result) >= Decimal("100000000000000"):
            raise ValueError()
        if result != result.quantize(Decimal("0.000001")) or (nonnegative and result < 0):
            raise ValueError()
        return result
    except (ValueError, InvalidOperation):
        raise HTTPException(422, detail="stocktake.quantity_invalid") from None


def parse(data):
    if len(data) > MAX_BYTES:
        raise HTTPException(413, detail="stocktake.file_too_large")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024:
                raise HTTPException(413, detail="stocktake.file_too_large")
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError, OSError, ParseError):
        raise HTTPException(422, detail="stocktake.file_invalid") from None
    try:
        sheet = wb.worksheets[0]
        rows = sheet.iter_rows()
        header = next(rows, ())
        header_names = tuple(str(c.value or "").strip() for c in header)
        if header_names not in (FIELDS, tuple(LABELS["th"][:7])):
            raise HTTPException(422, detail="stocktake.headers_invalid")
        result, seen = [], set()
        for number, cells in enumerate(rows, 2):
            if number > MAX_ROWS + 1:
                raise HTTPException(422, detail="stocktake.too_many_rows")
            if all(c.value is None for c in cells):
                continue
            if len(cells) != len(FIELDS) or any(c.data_type == "f" for c in cells):
                raise HTTPException(422, detail=f"stocktake.row_invalid:{number}")
            item = {
                key: str(c.value if c.value is not None else "").strip()
                for key, c in zip(FIELDS, cells)
            }
            if any(not item[k] for k in ("product_code", "product_name", "warehouse", "unit")):
                raise HTTPException(422, detail=f"stocktake.row_invalid:{number}")
            if any(len(v) > 300 for v in item.values()):
                raise HTTPException(422, detail=f"stocktake.row_invalid:{number}")
            # Identifiers must be text: numerical Excel cells can already have lost leading zeros.
            if any(cells[i].value is not None and cells[i].data_type != "s" for i in (0, 2)):
                raise HTTPException(422, detail=f"stocktake.code_as_text:{number}")
            item["book_qty"] = quantity(item["book_qty"])
            key = tuple(item[k] for k in ("product_code", "warehouse", "location"))
            if key in seen:
                raise HTTPException(422, detail=f"stocktake.duplicate:{number}")
            seen.add(key)
            result.append(item)
        if not result:
            raise HTTPException(422, detail="stocktake.empty")
        return result
    except (ParseError, IndexError, KeyError, ValueError):
        raise HTTPException(422, detail="stocktake.file_invalid") from None
    finally:
        wb.close()


def workbook(rows=None, lang="en"):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Stocktake"
    headers = LABELS["th"][:7]
    if rows is not None:
        headers = LABELS.get(lang, LABELS["en"])
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    if rows is None:
        for i, cell in enumerate(sheet[1]):
            cell.comment = Comment(" / ".join(values[i] for values in LABELS.values()), "Pearnly")
    for row in rows or []:
        values = [row[k] for k in FIELDS]
        actual = row["actual_qty"]
        values += [
            actual,
            None if actual is None else actual - row["book_qty"],
            str(row["counted_at"] or ""),
            str(row["counted_by"] or ""),
        ]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, Decimal) and len(cell.value.as_tuple().digits) > 15:
                cell.value = str(cell.value)
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for col in "ABCDEFG":
        sheet.column_dimensions[col].width = 24
    # Template identifier columns are preformatted as text.
    if rows is None:
        for row in sheet.iter_rows(min_row=2, max_row=1001, max_col=7):
            for cell in row:
                cell.number_format = "@" if cell.column != 7 else "0.######"
    sheet.auto_filter.ref = sheet.dimensions
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
