"""Tests for the DB-independent ingest core (parse -> normalize -> chunk)."""

import io
import unittest

from services.knowledge.ingest import (
    DEFAULT_MAX_CHARS,
    UnsupportedDocument,
    chunk_text,
    extract_text,
    ingest_document,
    normalize_text,
)

THAI = "ภาษีมูลค่าเพิ่มในประเทศไทยจัดเก็บที่ร้อยละเจ็ดของมูลค่าสินค้าหรือบริการ"


def _make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _make_xlsx(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_pdf(text: str) -> bytes:
    """Build a valid single-page PDF with a text layer (no extra dependency).

    Latin-only on purpose: it exercises the pypdf extraction plumbing. Thai PDFs
    depend on the source file's font/encoding and are validated with real client
    samples; image-only PDFs go to OCR.
    """
    stream = b"BT /F1 24 Tf 72 700 Td (" + text.encode("latin-1") + b") Tj ET"
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj" + obj + b"\nendobj\n"
    xref_pos = len(out)
    size = len(objs) + 1
    out += b"xref\n0 " + str(size).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer<</Size " + str(size).encode() + b"/Root 1 0 R>>\nstartxref\n"
    out += str(xref_pos).encode() + b"\n%%EOF"
    return bytes(out)


def test_normalize_crlf_trailing_and_blank_runs():
    raw = "line one  \r\nline two\r\n\r\n\r\n\r\npara two   "
    assert normalize_text(raw) == "line one\nline two\n\npara two"


def test_decode_thai_cp874():
    data = THAI.encode("cp874")
    assert extract_text("rules.txt", data) == THAI


def test_extract_text_utf8_and_bom():
    assert extract_text("a.md", "เก็บภาษี".encode("utf-8-sig")) == "เก็บภาษี"


def test_docx_extracts_paragraphs_thai():
    data = _make_docx([THAI, "บรรทัดที่สอง"])
    out = extract_text("policy.docx", data)
    assert THAI in out
    assert "บรรทัดที่สอง" in out


def test_xlsx_extracts_cells_thai():
    data = _make_xlsx([["ผู้ขาย", "ยอดเงิน"], ["สยามวัสดุ", 10000]])
    out = extract_text("ledger.xlsx", data)
    assert "ผู้ขาย" in out
    assert "สยามวัสดุ" in out
    assert "10000" in out


def test_pdf_extracts_text_layer():
    out = extract_text("invoice.pdf", _make_pdf("Invoice total 1500 THB"))
    assert "Invoice total 1500 THB" in out


def test_scanned_pdf_without_text_raises_for_ocr():
    with unittest.TestCase().assertRaises(UnsupportedDocument):
        extract_text("scan.pdf", _make_pdf(""))


def test_unknown_type_raises():
    with unittest.TestCase().assertRaises(UnsupportedDocument):
        extract_text("photo.png", b"\x89PNG")


def test_short_text_is_one_chunk():
    chunks = chunk_text("สั้นๆ พอ")
    assert len(chunks) == 1
    assert chunks[0].ordinal == 0
    assert chunks[0].char_count == len(chunks[0].text)


def test_paragraphs_pack_under_budget():
    text = "\n\n".join(["a" * 100, "b" * 100, "c" * 100])
    chunks = chunk_text(text, max_chars=250, overlap=0)
    # 100 + 1 + 100 = 201 fits; adding third (302) does not -> two chunks
    assert len(chunks) == 2
    assert [c.ordinal for c in chunks] == [0, 1]


def test_oversized_paragraph_window_splits_with_overlap():
    chunks = chunk_text("x" * 1000, max_chars=400, overlap=100)
    # step = 300 -> starts at 0,300,600,900 -> 4 windows
    assert len(chunks) == 4
    assert all(c.char_count <= 400 for c in chunks)
    assert chunks[1].text[:100] == chunks[0].text[-100:]  # overlap carried over


def test_chunk_text_rejects_nonpositive_budget():
    with unittest.TestCase().assertRaises(ValueError):
        chunk_text("anything", max_chars=0)


def test_ingest_document_end_to_end_thai():
    body = ("\r\n\r\n".join([THAI] * 30)).encode("utf-8")
    parsed = ingest_document("kb.txt", body, max_chars=DEFAULT_MAX_CHARS, overlap=80)
    assert parsed.filename == "kb.txt"
    assert parsed.chunks
    assert [c.ordinal for c in parsed.chunks] == list(range(len(parsed.chunks)))
    assert "\r" not in parsed.text


from tests.unit.knowledge._pytest_adapter import build_case  # noqa: E402

TestKnowledgeIngest = build_case(globals(), "TestKnowledgeIngest")
