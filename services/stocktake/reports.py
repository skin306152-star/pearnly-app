"""Thai product totals and count-location detail exported from one locked snapshot."""

from datetime import UTC
from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

SUMMARY = [
    "รหัสสินค้า",
    "ชื่อสินค้า",
    "บาร์โค้ด / รหัส QR",
    "หน่วย",
    "จำนวนตามบัญชี",
    "จำนวนที่นับได้รวม",
    "ผลต่าง",
    "ผลการตรวจนับ",
    "คลังสินค้าตามบัญชี",
    "คลังสินค้าที่ตรวจนับ",
    "ผลเทียบคลังสินค้า",
    "ตำแหน่งตามบัญชี",
    "ตำแหน่งที่ตรวจนับ",
    "ผลเทียบตำแหน่ง",
    "ดูรูปภาพ",
]
DETAIL = [
    "รหัสสินค้า",
    "ชื่อสินค้า",
    "บาร์โค้ด / รหัส QR",
    "คลังสินค้า",
    "ตำแหน่งจัดเก็บ",
    "หน่วย",
    "จำนวนที่นับครั้งนี้",
    "ผู้ตรวจนับ",
    "เวลาที่นับ (UTC)",
    "สถานะรายการ",
    "ผู้แก้ไขล่าสุด",
    "เวลาแก้ไข (UTC)",
    "คลังสินค้าตามบัญชี",
    "ผลเทียบคลังสินค้า",
    "ตำแหน่งตามบัญชี",
    "ผลเทียบตำแหน่ง",
    "ดูรูปภาพ",
]


def _comparison(book, actual_values):
    if not actual_values:
        return "ยังไม่ได้นับ"
    if not book:
        return "ไม่ได้ระบุในบัญชี"
    return "ตรงกัน" if all(value == book for value in actual_values) else "ไม่ตรงกัน"


def _append(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, Decimal) and len(cell.value.as_tuple().digits) > 15:
            cell.value = str(cell.value)
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _highlight(sheet, columns):
    for col in columns:
        sheet.cell(sheet.max_row, col).fill = PatternFill("solid", fgColor="FFF2CC")


def workbook(task):
    wb = Workbook()
    summary = wb.active
    summary.title = "สรุปผลตรวจนับ"
    _append(summary, SUMMARY)
    active = defaultdict(list)
    items = {row["id"]: row for row in task["items"]}
    for entry in task["entries"]:
        if not entry["voided"]:
            active[entry["item_id"]].append(entry)
    for row in task["items"]:
        warehouses = sorted({entry["warehouse"] for entry in active[row["id"]]})
        locations = sorted({entry["location"] for entry in active[row["id"]]})
        actual = row["actual_qty"]
        difference = None if actual is None else actual - row["book_qty"]
        result = (
            "ยังไม่ได้นับ"
            if actual is None
            else ("ตรงกัน" if difference == 0 else "เกินบัญชี" if difference > 0 else "ขาดบัญชี")
        )
        _append(
            summary,
            [
                row["product_code"],
                row["product_name"],
                row["barcode"],
                row["unit"],
                row["book_qty"],
                actual,
                difference,
                result,
                row["warehouse"],
                "\n".join(warehouses),
                _comparison(row["warehouse"], warehouses),
                row["location"],
                "\n".join(value or "(ไม่ระบุ)" for value in locations),
                _comparison(row["location"], locations),
                "",
            ],
        )
        if difference is not None and difference != 0:
            _highlight(summary, (5, 6, 7, 8))
        if _comparison(row["warehouse"], warehouses) == "ไม่ตรงกัน":
            _highlight(summary, (9, 10, 11))
        if _comparison(row["location"], locations) == "ไม่ตรงกัน":
            _highlight(summary, (12, 13, 14))
    detail = wb.create_sheet("รายละเอียดการนับ")
    _append(detail, DETAIL)
    for entry in task["entries"]:
        book = items[entry["item_id"]]
        _append(
            detail,
            [
                entry["product_code"],
                entry["product_name"],
                entry["barcode"],
                entry["warehouse"],
                entry["location"],
                entry["unit"],
                entry["quantity"],
                entry["counted_by_name"],
                entry["counted_at"].astimezone(UTC).isoformat(),
                "ยกเลิก (ไม่รวมยอด)" if entry["voided"] else "รวมยอด",
                entry["updated_by_name"],
                entry["updated_at"].astimezone(UTC).isoformat(),
                book["warehouse"],
                _comparison(book["warehouse"], [entry["warehouse"]]),
                book["location"],
                _comparison(book["location"], [entry["location"]]),
                "",
            ],
        )
        if not entry["voided"]:
            if _comparison(book["warehouse"], [entry["warehouse"]]) == "ไม่ตรงกัน":
                _highlight(detail, (4, 13, 14))
            if _comparison(book["location"], [entry["location"]]) == "ไม่ตรงกัน":
                _highlight(detail, (5, 15, 16))
    from services.stocktake.photo_report import attach

    attach(wb, task, summary, detail)
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="202033")
            cell.fill = PatternFill("solid", fgColor="EAF0FF")
            sheet.column_dimensions[cell.column_letter].width = 25
        sheet.column_dimensions["B"].width = 36
        sheet.row_dimensions[1].height = 26
        for cells in sheet:
            for cell in cells:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
